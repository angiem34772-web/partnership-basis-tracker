"""Write the rollforward results to a formatted Excel workbook."""

from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .loader import Inputs
from .rollforward import Results

MONEY = '#,##0;(#,##0);"-"'
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
FLAG_FILL = PatternFill("solid", fgColor="FCE4D6")

FLAG_DESCRIPTIONS = {
    "704D_LIMIT": "Loss limited to outside basis; excess suspended (IRC 704(d))",
    "SUSPENDED_LOSS_RELEASED": "Prior suspended losses allowed once basis was restored",
    "SUSPENDED_LOSS_FORFEITED": "Suspended losses lost when the interest was liquidated",
    "731_GAIN": "Cash or deemed distributions exceeded basis; gain recognized (IRC 731(a)(1))",
    "NEGATIVE_TAX_CAPITAL": "Tax-basis capital account went negative",
    "NONDEDUCTIBLE_OVER_BASIS": "Nondeductible expenses exceeded remaining basis",
    "INSTALLMENT_BUYOUT": "Buyout paid in installments (IRC 736(b), Reg. 1.736-1(b)(6))",
    "LIQUIDATION_GAIN": "Gain on buyout payments",
    "LIQUIDATION_LOSS": "Loss on final buyout payment (IRC 731(a)(2))",
    "RECONCILIATION": "Basis does not reconcile to capital plus liabilities",
}


def export(inputs: Inputs, results: Results, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    names = {p.partner_id: p.name for p in inputs.partners}
    entity_names = {e.entity_id: e.name for e in inputs.entities}

    wb = Workbook()
    _summary(wb.active, inputs, results, entity_names)
    _partner_summary(wb.create_sheet("Partner Summary"), inputs, results, names)
    _basis(wb.create_sheet("Basis Rollforward"), results, names)
    _capital(wb.create_sheet("Capital Accounts"), results, names)
    _flags(wb.create_sheet("Flags"), results, names)
    _buyouts(wb.create_sheet("Buyout Schedules"), inputs, results, names)
    _division(wb.create_sheet("2022 Division"), results, names)
    _k1_inputs(wb.create_sheet("K-1 Inputs"), inputs, names)
    wb.save(path)
    return path


def _table(ws, headers, rows, money_from: int, widths=None, flag_col: int | None = None):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 45
    for row in rows:
        ws.append([round(v, 2) if isinstance(v, float) else v for v in row])
    for row in ws.iter_rows(min_row=2, min_col=money_from):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = MONEY
    if flag_col:
        for row in ws.iter_rows(min_row=2):
            if row[flag_col - 1].value:
                for cell in row:
                    cell.fill = FLAG_FILL
    for i, header in enumerate(headers, start=1):
        width = (widths or {}).get(i, max(12, min(len(str(header)) + 2, 18)))
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "E2" if money_from > 4 else "B2"
    ws.auto_filter.ref = ws.dimensions


def _summary(ws, inputs, results, entity_names):
    ws.title = "Summary"
    ws["A1"] = "Partnership Basis Tracker"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = "Sample data only. All partners, partnerships, and amounts are fictional."
    ws["A2"].font = Font(italic=True, color="7F7F7F")

    years = [y.tax_year for y in results.years]
    lines = [
        ("Partnerships", ", ".join(f"{e.name} ({e.entity_id})" for e in inputs.entities)),
        ("Partners tracked", len(inputs.partners)),
        ("Years covered", f"{min(years)} - {max(years)}"),
        ("Partner-years rolled forward", len(results.years)),
        ("Buyout payments scheduled", len(results.payments)),
        ("Reconciliation exceptions", sum(1 for y in results.years if abs(y.unexplained_difference) > 1)),
    ]
    row = 4
    for label, value in lines:
        ws.cell(row, 1, label).font = Font(bold=True)
        ws.cell(row, 2, value)
        row += 1

    row += 1
    ws.cell(row, 1, "Flags raised").font = Font(bold=True, size=12)
    row += 1
    for code, count in sorted(Counter(f.code for f in results.flags).items()):
        ws.cell(row, 1, code)
        ws.cell(row, 2, count)
        ws.cell(row, 3, FLAG_DESCRIPTIONS.get(code, ""))
        row += 1

    row += 1
    ws.cell(row, 1, "How to read this workbook").font = Font(bold=True, size=12)
    guide = [
        ("Partner Summary", "One line per partner per partnership: where each partner stands today."),
        ("Basis Rollforward", "Outside basis, year by year, in the IRC 705 / Reg. 1.704-1(d)(2) order."),
        ("Capital Accounts", "Tax-basis capital, liability share, and a check that basis ties out."),
        ("Flags", "Every basis limitation, distribution gain, and buyout event, with the amount."),
        ("Buyout Schedules", "Each buyout payment, showing basis recovered first, then gain or loss."),
        ("2022 Division", "How each partner's basis, capital, and debt split between MHP and MHD."),
        ("K-1 Inputs", "The K-1 amounts the tracker started from."),
    ]
    for sheet, text in guide:
        row += 1
        ws.cell(row, 1, sheet).font = Font(bold=True)
        ws.cell(row, 2, text)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 70


def _partner_summary(ws, inputs, results, names):
    by_key = defaultdict(list)
    for y in results.years:
        by_key[(y.partner_id, y.entity_id)].append(y)
    admitted = {p.partner_id: p for p in inputs.partners}
    rows = []
    entity_order = [e.entity_id for e in inputs.entities]
    for (pid, eid), ys in sorted(by_key.items(), key=lambda kv: (kv[0][0], entity_order.index(kv[0][1]))):
        last = ys[-1]
        rows.append([
            pid, names[pid], eid, admitted[pid].admission_type, f"{ys[0].tax_year}-{last.tax_year}", last.status,
            sum(y.contributions for y in ys),
            sum(y.income_items + y.tax_exempt_income for y in ys),
            sum(y.losses_allowed for y in ys),
            sum(y.cash_distributions for y in ys),
            sum(y.gain_on_distributions for y in ys),
            sum(y.liquidating_payments for y in ys),
            sum(y.liquidation_gain - y.liquidation_loss for y in ys),
            last.ending_basis, last.suspended_losses_end, last.ending_capital, last.ending_liabilities,
        ])
    headers = [
        "Partner", "Name", "Entity", "Admission", "Years", "Status",
        "Total contributions", "Total income", "Total losses allowed", "Total cash distributions",
        "Total 731 gain", "Buyout payments", "Buyout gain / (loss)",
        "Ending basis", "Suspended losses", "Ending tax capital", "Ending liability share",
    ]
    _table(ws, headers, rows, money_from=7, widths={2: 28, 6: 26})


def _basis(ws, results, names):
    headers = [
        "Partner", "Name", "Entity", "Year", "Beginning basis", "2022 division", "Contributions",
        "Income items", "Tax-exempt income", "Increase in liability share", "Cash distributions",
        "Decrease in liability share (deemed dist.)", "731 gain", "Nondeductible expenses",
        "Current-year losses", "Prior suspended losses", "Losses allowed", "Suspended losses c/f",
        "Suspended losses forfeited", "Buyout payments", "Buyout gain", "Buyout loss", "Ending basis",
        "Status", "Flags",
    ]
    rows = [
        [
            y.partner_id, names[y.partner_id], y.entity_id, y.tax_year, y.beginning_basis, y.division_adjustment,
            y.contributions, y.income_items, y.tax_exempt_income, y.liability_increase, -y.cash_distributions,
            -y.liability_decrease, y.gain_on_distributions, -y.nondeductible_allowed, y.current_losses,
            y.prior_suspended_losses, -y.losses_allowed, y.suspended_losses_end, y.suspended_losses_forfeited,
            -y.liquidating_payments, y.liquidation_gain, y.liquidation_loss, y.ending_basis, y.status,
            ", ".join(f for f in y.flags if f != "NEGATIVE_TAX_CAPITAL"),
        ]
        for y in results.years
    ]
    _table(ws, headers, rows, money_from=5, widths={2: 26, 24: 24, 25: 34}, flag_col=25)
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = "0"


def _capital(ws, results, names):
    headers = [
        "Partner", "Name", "Entity", "Year", "Beginning tax capital", "2022 division", "Contributions",
        "Net income / (loss)", "Nondeductible expenses", "Cash distributions", "Closed on buyout",
        "Ending tax capital", "Ending liability share", "Ending outside basis",
        "Basis less (capital + liabilities)", "Explained by suspended losses, 731 gains, buyouts",
        "Unexplained difference",
    ]
    rows = [
        [
            y.partner_id, names[y.partner_id], y.entity_id, y.tax_year, y.beginning_capital,
            y.capital_division_adjustment, y.contributions, y.net_income_for_capital,
            -y.nondeductible_expenses, -y.cash_distributions, y.capital_closed_on_liquidation, y.ending_capital,
            y.ending_liabilities, y.ending_basis, y.basis_capital_difference, y.explained_difference,
            y.unexplained_difference,
        ]
        for y in results.years
    ]
    _table(ws, headers, rows, money_from=5, widths={2: 26, 16: 22})
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = "0"


def _flags(ws, results, names):
    headers = ["Partner", "Name", "Entity", "Year", "Flag", "Amount", "Detail"]
    rows = [
        [f.partner_id, names[f.partner_id], f.entity_id, f.tax_year, f.code, f.amount, f.message]
        for f in sorted(results.flags, key=lambda f: (f.tax_year, f.entity_id, f.partner_id))
    ]
    _table(ws, headers, rows, money_from=6, widths={2: 26, 5: 26, 6: 14, 7: 110})
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = "0"


def _buyouts(ws, inputs, results, names):
    headers = [
        "Partner", "Name", "Entity", "Payment date", "Payment #", "Of", "Payment", "Basis before",
        "Basis recovered", "Gain", "Loss", "Basis after", "Method",
    ]
    method = {(b.partner_id, b.entity_id): b.method for b in inputs.buyouts}
    rows = [
        [
            r.payment.partner_id, names[r.payment.partner_id], r.payment.entity_id, r.payment.paid_on,
            r.payment.number, r.payment.of, r.payment.amount, r.basis_before, r.basis_recovered, r.gain,
            r.loss, r.basis_after, method[(r.payment.partner_id, r.payment.entity_id)],
        ]
        for r in sorted(results.payments, key=lambda r: (r.payment.partner_id, r.payment.entity_id, r.payment.paid_on))
    ]
    _table(ws, headers, rows, money_from=7, widths={2: 26, 4: 14, 13: 14})
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = "mm/dd/yyyy"


def _division(ws, results, names):
    headers = [
        "Partner", "Name", "From", "To", "Basis before division", "Basis moved to new partnership",
        "Basis kept", "Tax capital moved", "Liability share moved", "Suspended losses moved",
    ]
    rows = [
        [
            d.partner_id, names[d.partner_id], d.from_entity, d.to_entity, d.basis_before, d.basis_transferred,
            d.basis_before - d.basis_transferred, d.capital_transferred, d.liabilities_transferred,
            d.suspended_transferred,
        ]
        for d in results.divisions
    ]
    _table(ws, headers, rows, money_from=5, widths={2: 26})


def _k1_inputs(ws, inputs, names):
    headers = [
        "Partner", "Name", "Entity", "Year", "Contributions", "Ordinary income / (loss)", "Interest income",
        "Net capital gain / (loss)", "Tax-exempt income", "Nondeductible expenses", "Liability share, end of year",
    ]
    rows = [
        [
            k.partner_id, names[k.partner_id], k.entity_id, k.tax_year, k.contributions, k.ordinary_income,
            k.interest_income, k.net_capital_gain, k.tax_exempt_income, k.nondeductible_expenses,
            k.liability_share_end,
        ]
        for k in sorted(inputs.k1s, key=lambda k: (k.partner_id, k.entity_id, k.tax_year))
    ]
    _table(ws, headers, rows, money_from=5, widths={2: 26})
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = "0"
